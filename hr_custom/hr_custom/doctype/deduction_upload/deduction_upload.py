# Copyright (c) 2026, wangui and contributors
# For license information, please see license.txt

import os

import frappe
from frappe.model.document import Document
from frappe.utils import cstr, flt, get_last_day, getdate
from frappe.utils.csvutils import read_csv_content
from frappe.utils.xlsxutils import read_xlsx_file_from_attached_file

from hr_custom.hr_custom.salary_slip_sync import deferred_slip_refresh, get_salary_slips

EMPLOYEE_ID_COL = "Employee ID"
EMPLOYEE_NAME_COL = "Employee Name"
REMARKS_COL = "Remarks"


class DeductionUpload(Document):
	def validate(self):
		self.payroll_date = get_last_day(self.payroll_month)
		self.validate_items()
		self.row_count = len(self.items)
		self.total_amount = sum(flt(d.amount) for d in self.items)

	def validate_items(self):
		components = set(self.get_component_names())
		seen, errors = set(), []
		for d in self.items:
			if d.salary_component not in components:
				errors.append(f"Row {d.idx}: {d.salary_component} is not one of the selected deduction columns")
			if flt(d.amount) <= 0:
				errors.append(f"Row {d.idx}: amount must be greater than zero")
			key = (d.employee, d.salary_component)
			if key in seen:
				errors.append(f"Row {d.idx}: {d.employee} has {d.salary_component} more than once")
			seen.add(key)
		if errors:
			frappe.throw("<br>".join(errors), title="Invalid Rows")

	def before_submit(self):
		if not self.items:
			frappe.throw("Nothing to submit. Attach the filled template and click Load File first.")
		self.validate_employees_payable()
		self.validate_no_submitted_salary_slips()

	def validate_employees_payable(self):
		"""Additional Salary rejects inactive employees and those without a salary structure, so report them all up front."""
		employees = list({d.employee for d in self.items})
		active = set(frappe.get_all("Employee", filters={"name": ["in", employees], "status": "Active"}, pluck="name"))
		with_structure = set(frappe.get_all(
			"Salary Structure Assignment",
			filters={"employee": ["in", employees], "docstatus": 1, "from_date": ["<=", self.payroll_date]},
			pluck="employee",
		))
		errors = []
		for d in self.items:
			if d.employee not in active:
				errors.append(f"Row {d.idx}: {d.employee} ({d.employee_name}) is not an active employee")
			elif d.employee not in with_structure:
				errors.append(f"Row {d.idx}: {d.employee} ({d.employee_name}) has no Salary Structure Assignment on or before {self.payroll_date}")
		if errors:
			frappe.throw("<br>".join(errors), title="Employees Cannot Be Paid")

	def validate_no_submitted_salary_slips(self):
		"""A deduction dated in a month whose slip is already submitted would never be paid, so stop here."""
		errors = []
		for employee in sorted({d.employee for d in self.items}):
			slips = get_salary_slips(employee, self.payroll_date, 1)
			if slips:
				errors.append(f"{employee}: {', '.join(slips)}")
		if errors:
			frappe.throw(
				"Salary slips for this month are already submitted for:<br>" + "<br>".join(errors)
				+ "<br><br>Either cancel and amend those slips first, or change the Payroll Month to next month "
				"so the deductions go into the next payroll.",
				title="Payroll Already Submitted",
			)

	def on_submit(self):
		# Draft salary slips from an existing Payroll Entry are re-saved once per employee at the end
		with deferred_slip_refresh():
			self.create_additional_salaries()

	def create_additional_salaries(self):
		for d in self.items:
			existing = frappe.get_all(
				"Additional Salary",
				filters={
					"employee": d.employee,
					"salary_component": d.salary_component,
					"payroll_date": self.payroll_date,
					"docstatus": ["!=", 2],
				},
				pluck="name",
			)
			if existing and self.if_exists == "Skip":
				d.db_set("status", f"Skipped - exists ({existing[0]})")
				continue
			for name in existing:
				existing_doc = frappe.get_doc("Additional Salary", name)
				if existing_doc.docstatus == 1:
					existing_doc.cancel()
				else:
					existing_doc.delete()

			try:
				add_sal = frappe.new_doc("Additional Salary")
				add_sal.employee = d.employee
				add_sal.salary_component = d.salary_component
				add_sal.type = "Deduction"
				add_sal.amount = d.amount
				add_sal.payroll_date = self.payroll_date
				add_sal.company = self.company
				add_sal.ref_doctype = self.doctype
				add_sal.ref_docname = self.name
				add_sal.overwrite_salary_structure_amount = 0
				add_sal.insert()
				add_sal.submit()
			except Exception as e:
				frappe.throw(f"Row {d.idx} ({d.employee}, {d.salary_component}): {e}")

			d.db_set({"additional_salary": add_sal.name, "status": "Replaced" if existing else "Created"})

	def on_cancel(self):
		self.ignore_linked_doctypes = ("Additional Salary",)
		with deferred_slip_refresh():
			self.cancel_additional_salaries()

	def cancel_additional_salaries(self):
		for d in self.items:
			if not d.additional_salary:
				continue
			add_sal = frappe.get_doc("Additional Salary", d.additional_salary)
			if add_sal.docstatus == 1:
				add_sal.cancel()
			d.db_set("status", "Cancelled")

	def get_component_names(self):
		return [c.salary_component for c in self.components]

	@frappe.whitelist()
	def load_from_file(self):
		"""Read the attached template into the items table, replacing any rows already loaded."""
		if not self.upload_file:
			frappe.throw("Attach the filled template first.")

		rows = read_uploaded_file(self.upload_file)
		self.set("items", parse_template_rows(rows, self.get_component_names()))
		self.save()
		frappe.msgprint(f"Loaded {len(self.items)} deductions totalling {frappe.format(self.total_amount, 'Currency')}.")


def read_uploaded_file(file_url):
	file_doc = frappe.get_doc("File", {"file_url": file_url})
	extension = os.path.splitext(file_doc.file_name or file_url)[1].lower()
	if extension == ".xlsx":
		return read_xlsx_file_from_attached_file(file_url=file_url)
	if extension == ".csv":
		return read_csv_content(file_doc.get_content())
	frappe.throw("Upload the template as .xlsx or .csv")


def parse_template_rows(rows, components):
	"""Turn the wide template (one column per component) into one item per employee per component."""
	rows = [r for r in rows if any(cstr(c).strip() for c in r)]
	if not rows:
		frappe.throw("The uploaded file is empty.")

	header = [cstr(c).strip().lower() for c in rows[0]]
	col = {name: i for i, name in enumerate(header) if name}

	emp_idx = col.get(EMPLOYEE_ID_COL.lower(), col.get("employee"))
	name_idx = col.get(EMPLOYEE_NAME_COL.lower())
	remarks_idx = col.get(REMARKS_COL.lower())
	if emp_idx is None and name_idx is None:
		frappe.throw(f"The file needs an '{EMPLOYEE_ID_COL}' or '{EMPLOYEE_NAME_COL}' column. Use the downloaded template.")

	component_cols = []
	for c in components:
		# Match "SACCO Contribution - Amount (KES)" from the template as well as a plain "SACCO Contribution" header
		matches = [i for name, i in col.items() if name == c.lower() or name.startswith(f"{c.lower()} - amount")]
		if matches:
			component_cols.append((c, matches[0]))
	if not component_cols:
		frappe.throw(f"None of the selected deduction columns ({', '.join(components)}) were found in the file header.")

	employees = frappe.get_all("Employee", fields=["name", "employee_name"])
	valid_ids = {e.name for e in employees}
	by_name = {}
	for e in employees:
		by_name.setdefault(cstr(e.employee_name).strip().lower(), []).append(e.name)

	cell = lambda row, idx: cstr(row[idx]).strip() if idx is not None and idx < len(row) and row[idx] is not None else ""

	items, errors = [], []
	for line_no, row in enumerate(rows[1:], start=2):
		amounts = [(c, flt(cell(row, i))) for c, i in component_cols if cell(row, i)]
		if not any(amount for _, amount in amounts):
			continue

		emp_id, emp_name = cell(row, emp_idx), cell(row, name_idx)
		if emp_id:
			if emp_id not in valid_ids:
				errors.append(f"Line {line_no}: employee ID '{emp_id}' not found")
				continue
			employee = emp_id
		else:
			matches = by_name.get(emp_name.lower(), [])
			if len(matches) != 1:
				reason = "not found" if not matches else f"matches {len(matches)} employees, fill in the Employee ID"
				errors.append(f"Line {line_no}: employee name '{emp_name}' {reason}")
				continue
			employee = matches[0]

		for component, amount in amounts:
			if amount < 0:
				errors.append(f"Line {line_no}: {component} amount cannot be negative")
			elif amount > 0:
				items.append({
					"employee": employee,
					"salary_component": component,
					"amount": amount,
					"remarks": cell(row, remarks_idx),
				})

	if errors:
		frappe.throw("<br>".join(errors), title="Could not load file")
	if not items:
		frappe.throw("No amounts found in the file.")
	return items


@frappe.whitelist()
def download_template(company, components, payroll_month=None):
	frappe.has_permission("Deduction Upload", "read", throw=True)
	components = frappe.parse_json(components) if isinstance(components, str) else components
	if not components:
		frappe.throw("Add at least one row to Deduction Columns first - each one becomes an amount column.")

	employees = frappe.get_all(
		"Employee",
		filters={"status": "Active", "company": company},
		fields=["name", "employee_name", "department"],
		order_by="employee_name",
	)
	month = getdate(payroll_month).strftime("%B %Y") if payroll_month else ""
	currency = frappe.get_cached_value("Company", company, "default_currency") or ""
	content = build_template_workbook(employees, components, company, month, currency)

	from frappe.desk.utils import provide_binary_file

	provide_binary_file(f"Deduction Upload {month or 'Template'}", "xlsx", content)


def amount_header(component, currency):
	return f"{component} - Amount ({currency})" if currency else f"{component} - Amount"


def build_template_workbook(employees, components, company, month, currency):
	from io import BytesIO

	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter
	from openpyxl.worksheet.datavalidation import DataValidation

	header_fill = PatternFill("solid", fgColor="1F4E78")
	locked_fill = PatternFill("solid", fgColor="F2F2F2")
	input_fill = PatternFill("solid", fgColor="FFF2CC")
	thin = Side(style="thin", color="BFBFBF")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)

	wb = Workbook()
	ws = wb.active
	ws.title = "Deductions"

	headers = [EMPLOYEE_ID_COL, EMPLOYEE_NAME_COL, *[amount_header(c, currency) for c in components], REMARKS_COL]
	ws.append(headers)
	for cell in ws[1]:
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = header_fill
		cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
		cell.border = border
	ws.row_dimensions[1].height = 36

	first_amount_col, last_amount_col = 3, 2 + len(components)
	remarks_col = last_amount_col + 1
	for e in employees:
		ws.append([e.name, e.employee_name, *([None] * len(components)), None])
		row = ws.max_row
		for col in range(1, remarks_col + 1):
			cell = ws.cell(row=row, column=col)
			cell.border = border
			if col < first_amount_col:
				cell.fill = locked_fill
			elif col <= last_amount_col:
				cell.fill = input_fill
				cell.number_format = "#,##0.00"

	# Reject text and negative numbers in the amount cells
	last_row = max(ws.max_row, 2) + 200
	validation = DataValidation(
		type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
		showErrorMessage=True, errorTitle="Invalid amount", error="Enter a number (0 or more), or leave blank.",
	)
	validation.add(f"{get_column_letter(first_amount_col)}2:{get_column_letter(last_amount_col)}{last_row}")
	ws.add_data_validation(validation)

	ws.column_dimensions["A"].width = 18
	ws.column_dimensions["B"].width = 32
	for col in range(first_amount_col, last_amount_col + 1):
		ws.column_dimensions[get_column_letter(col)].width = 22
	ws.column_dimensions[get_column_letter(remarks_col)].width = 30
	ws.freeze_panes = "C2"

	info = wb.create_sheet("Instructions")
	lines = [
		("Deduction Upload Template", True),
		(f"Company: {company}", False),
		(f"Payroll month: {month or '(set on the Deduction Upload form)'}", False),
		("", False),
		("How to fill it in", True),
		("1. Go to the 'Deductions' sheet. Each row is one active employee.", False),
		("2. Type the amount to deduct in the yellow cells - one column per deduction type:", False),
		*[(f"      - {c}", False) for c in components],
		("3. Leave a cell blank (or 0) if the employee has no such deduction this month.", False),
		("4. Employees not listed: add a new row with their Employee ID (or exact full name).", False),
		("5. Do not rename or delete the header row. You may delete rows you don't need.", False),
		("6. Save as .xlsx and attach it to the 'Filled Template' field on the Deduction Upload form.", False),
		("   The rows load automatically - check them, then Submit.", False),
	]
	for text, bold in lines:
		info.append([text])
		if bold:
			info.cell(row=info.max_row, column=1).font = Font(bold=True, size=12 if info.max_row > 1 else 14)
	info.column_dimensions["A"].width = 100

	wb.active = 0
	out = BytesIO()
	wb.save(out)
	return out.getvalue()


@frappe.whitelist()
def get_default_components():
	"""Deduction components that look like SACCO, loan or salary advance deductions, used to prefill a new upload."""
	return frappe.get_all(
		"Salary Component",
		filters={"type": "Deduction", "disabled": 0},
		or_filters=[["name", "like", "%sacco%"], ["name", "like", "%loan%"], ["name", "like", "%advance%"]],
		pluck="name",
		order_by="name",
	)
